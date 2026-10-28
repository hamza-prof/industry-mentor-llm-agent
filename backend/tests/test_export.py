import io

from docx import Document as Docx
from docx.oxml import parse_xml

from app.export import exported, final_text, to_docx, to_markdown
from app.models import Review, ReviewComment
from app.schemas import (
    Figure,
    Paragraph,
    ParsedDocument,
    Section,
    StructuredSummary,
    SummaryField,
)
from tests.helpers import add_picture, make_png

FIELDS = (
    "research_problem",
    "objectives",
    "technical_approach",
    "main_conclusions",
    "innovations",
    "limitations",
)


def parsed(language="en") -> ParsedDocument:
    fig = Figure(image_id="a" * 32, file=f"{'a' * 32}.png", source="embedded", caption="Fig 1 x")
    return ParsedDocument(
        filename="u.md",
        doc_type="report",
        language=language,
        title="Order system report",
        sections=[
            Section(
                id="S1",
                heading="1 Approach",
                level=1,
                paragraphs=[Paragraph(id="S1.P1", text="We moved to microservices.")],
            ),
            Section(
                id="S2",
                heading="2 Results",
                level=1,
                paragraphs=[
                    Paragraph(id="S2.P1", text="Users are happier now."),
                    Paragraph(id="S2.P2", text="[Figure] Fig 1 x", kind="figure", figure=fig),
                ],
            ),
        ],
    )


def review(**kw) -> Review:
    summary = StructuredSummary(
        doc_type="report", **{f: SummaryField(text=f"{f} text") for f in FIELDS}
    )
    return Review(
        id=1,
        document_id=1,
        status="done",
        model="m-1",
        summary=summary.model_dump(mode="json"),
        overall_feedback=kw.pop("overall", "Solid work overall."),
        **kw,
    )


def comment(status="accepted", pid="S2.P1", origin="ai", **kw) -> ReviewComment:
    sid = pid.split(".")[0] if pid else kw.pop("sid", None)
    return ReviewComment(
        review_id=1,
        section_id=sid,
        paragraph_id=pid,
        anchored=pid is not None,
        category=kw.pop("category", "evidence"),
        severity=kw.pop("severity", "high"),
        problem=kw.pop("problem", f"Problem at {pid}"),
        suggestion=kw.pop("suggestion", "Add data."),
        revision_steps=kw.pop("steps", ["Collect a survey"]),
        origin=origin,
        status=status,
        **kw,
    )


def test_exported_keeps_only_confirmed():
    cs = [
        comment("accepted"),
        comment("edited", mentor_text="mine"),
        comment("rejected"),
        comment("pending"),
        comment("accepted", origin="mentor"),
    ]
    assert [c.status for c in exported(cs)] == ["accepted", "edited", "accepted"]


def test_final_text_uses_mentor_text_when_edited():
    assert final_text(comment("edited", mentor_text="Show the survey.")) == "Show the survey."
    assert "Problem at S2.P1" in final_text(comment()) and "Add data." in final_text(comment())


def test_markdown_has_header_overall_summary_and_grouped_comments():
    cs = [
        comment(pid="S2.P1", problem="Results claim lacks data"),
        comment(pid="S1.P1", problem="Approach unclear"),
        comment("rejected", problem="Rejected one"),
        comment("edited", pid="S1.P1", mentor_text="Mentor rewrote this"),
    ]
    md = to_markdown(review(), "报告.md", parsed(), cs)
    assert "报告.md" in md and "m-1" in md
    assert "Solid work overall." in md and "research_problem text" in md
    assert "Rejected one" not in md and "Mentor rewrote this" in md
    assert md.index("1 Approach") < md.index("Approach unclear") < md.index("2 Results")
    assert md.index("2 Results") < md.index("Results claim lacks data")
    assert "We moved to microservices." in md  # paragraph excerpt shown
    assert "Collect a survey" in md


def test_markdown_document_level_comment_and_empty_note():
    md = to_markdown(review(), "r.md", parsed(), [comment(pid=None, problem="No references")])
    assert "No references" in md
    md = to_markdown(review(overall=None), "r.md", parsed(), [comment("rejected")])
    assert "No comments were confirmed" in md


def test_markdown_in_chinese_for_chinese_document():
    md = to_markdown(review(), "r.md", parsed("zh"), [comment()])
    assert "总体评价" in md and "修改步骤" in md


def test_new_docx_for_non_docx_source(tmp_path):
    (tmp_path / f"{'a' * 32}.png").write_bytes(make_png())
    cs = [comment(), comment(pid="S2.P2", problem="Axis has no unit"), comment("rejected")]
    data = to_docx(review(), "r.pdf", parsed(), cs, original=None, figures_dir=tmp_path)
    d = Docx(io.BytesIO(data))
    text = "\n".join(p.text for p in d.paragraphs)
    assert "Solid work overall." in text and "research_problem text" in text
    [table] = d.tables
    assert len(table.rows) == 1 + 2  # header + exported comments
    assert "Axis has no unit" in table.rows[2].cells[2].text
    assert len(d.inline_shapes) == 1  # figure thumbnail embedded with its comment


HYPERLINK = (
    '<w:hyperlink xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
    "<w:r><w:t>Users are happier now.</w:t></w:r></w:hyperlink>"
)


def original_docx(tmp_path, runless=False):
    d = Docx()
    d.add_heading("1 Approach", 1)
    d.add_paragraph("We moved to microservices.")
    d.add_heading("2 Results", 1)
    if runless:  # text only inside a hyperlink → paragraph.runs is empty
        d.add_paragraph()._p.append(parse_xml(HYPERLINK))
        d.add_paragraph("Next paragraph.")
    else:
        d.add_paragraph("Users are happier now.")
    add_picture(d.add_paragraph(), make_png())
    d.add_paragraph("Fig 1 x")
    path = tmp_path / "orig.docx"
    d.save(path)
    return path


def test_docx_source_gets_word_comments_on_anchored_paragraphs(tmp_path):
    cs = [
        comment(pid="S1.P1", problem="Approach unclear"),
        comment("edited", pid="S2.P1", mentor_text="Cite the survey"),
        comment(pid="S2.P2", problem="Axis has no unit"),
        comment(pid=None, problem="No references section"),
        comment("rejected", pid="S1.P1"),
    ]
    data = to_docx(review(), "r.docx", parsed(), cs, original_docx(tmp_path), tmp_path)
    d = Docx(io.BytesIO(data))
    assert len(list(d.comments)) == 3  # anchored, exported
    texts = {c.text for c in d.comments}
    assert any("Approach unclear" in t for t in texts) and "Cite the survey" in texts
    assert "We moved to microservices." in [p.text for p in d.paragraphs]  # original kept
    tail = "\n".join(p.text for p in d.paragraphs[-6:])
    assert "No references section" in tail and "Solid work overall." in "\n".join(
        p.text for p in d.paragraphs
    )


def test_comment_lands_on_the_matching_paragraph(tmp_path):
    cs = [comment("edited", pid="S2.P1", mentor_text="Cite the survey")]
    data = to_docx(review(), "r.docx", parsed(), cs, original_docx(tmp_path), tmp_path)
    xml = Docx(io.BytesIO(data)).element.body.xml
    start = xml.index("commentRangeStart")
    assert xml.index("Users are happier now.") > start > xml.index("2 Results")


def test_paragraph_without_runs_uses_nearest_paragraph(tmp_path):
    cs = [comment("edited", pid="S2.P1", mentor_text="Cite the survey")]
    data = to_docx(review(), "r.docx", parsed(), cs, original_docx(tmp_path, True), tmp_path)
    d = Docx(io.BytesIO(data))
    assert len(list(d.comments)) == 1
    xml = d.element.body.xml
    assert xml.index("Next paragraph.") > xml.index("commentRangeStart") > xml.index("happier")


def test_unmatched_paragraph_falls_back_to_section_heading(tmp_path):
    p = parsed()
    p.sections[1].paragraphs[0].text = "Text the Word file does not contain."
    cs = [comment(pid="S2.P1", problem="Lost")]
    d = Docx(io.BytesIO(to_docx(review(), "r.docx", p, cs, original_docx(tmp_path), tmp_path)))
    assert len(list(d.comments)) == 1
    xml = d.element.body.xml
    assert xml.index("2 Results") > xml.index("commentRangeStart") > xml.index("microservices")


def test_figure_comment_lands_on_the_image_itself(tmp_path):
    import hashlib

    p = parsed()
    fig = p.sections[1].paragraphs[1].figure
    fig.image_id, fig.caption = hashlib.sha256(make_png()).hexdigest()[:32], None
    cs = [comment(pid="S2.P2", problem="Axis has no unit")]
    d = Docx(io.BytesIO(to_docx(review(), "r.docx", p, cs, original_docx(tmp_path), tmp_path)))
    xml = d.element.body.xml
    assert xml.index("<w:drawing") > xml.index("commentRangeStart") > xml.index("happier")
