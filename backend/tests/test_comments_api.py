import io
from urllib.parse import quote

import pytest
from docx import Document as Docx
from langchain_core.messages import AIMessage

from app.runner import ReviewModels
from tests.test_reviews_api import set_status, upload
from tests.test_runner import FakeChat


class OverallChat(FakeChat):
    def __init__(self):
        super().__init__()
        self.prompts = []

    async def ainvoke(self, msgs):
        self.prompts.append("\n".join(m.content for m in msgs))
        return AIMessage("  Good structure; add evidence to the results.  ")


@pytest.fixture
def chat():
    return OverallChat()


@pytest.fixture
def app_client(make_client, chat):
    with make_client(llm_api_key="k") as c:
        c.app.state.review_models = ReviewModels(chat=chat)
        yield c


def reviewed(client, **kw):
    """Upload + run a (fake) review → (review id, its comments)."""
    rid = client.post("/api/reviews", json={"document_id": upload(client, **kw)}).json()["id"]
    return rid, client.get(f"/api/reviews/{rid}").json()["comments"]


def ai(comments):
    return [c for c in comments if c["origin"] == "ai"]


def patch(client, cid, **body):
    return client.patch(f"/api/comments/{cid}", json=body)


def test_accept_reject_edit_and_undo(app_client):
    _, cs = reviewed(app_client)
    cid = ai(cs)[0]["id"]
    assert patch(app_client, cid, status="accepted").json()["status"] == "accepted"
    assert patch(app_client, cid, status="rejected").json()["status"] == "rejected"
    r = patch(app_client, cid, status="edited", mentor_text="Show the survey data.")
    assert r.json()["status"] == "edited" and r.json()["mentor_text"] == "Show the survey data."
    assert patch(app_client, cid, status="pending").json()["status"] == "pending"
    r = patch(app_client, cid, severity="low", category="clarity")
    assert (r.json()["severity"], r.json()["category"]) == ("low", "clarity")


def test_edited_needs_mentor_text(app_client):
    _, cs = reviewed(app_client)
    r = patch(app_client, ai(cs)[0]["id"], status="edited", mentor_text="   ")
    assert r.status_code == 422 and r.json()["code"] == "EMPTY_EDIT"
    assert patch(app_client, ai(cs)[0]["id"], status="edited").status_code == 422


def test_patch_unknown_comment_404_and_bad_status_422(app_client):
    assert patch(app_client, 999, status="accepted").status_code == 404
    _, cs = reviewed(app_client)
    assert patch(app_client, ai(cs)[0]["id"], status="done").status_code == 422


def test_mentor_adds_comment(app_client):
    rid, _ = reviewed(app_client)
    r = app_client.post(
        f"/api/reviews/{rid}/comments",
        json={"paragraph_id": "S2.P1", "quote": "Users are happier", "problem": "Which users?"},
    )
    assert r.status_code == 201, r.text
    c = r.json()
    assert (c["origin"], c["status"], c["section_id"], c["anchored"]) == (
        "mentor",
        "accepted",
        "S2",
        True,
    )
    assert c["quote"] == "Users are happier"
    got = app_client.get(f"/api/reviews/{rid}").json()["comments"]
    assert any(x["id"] == c["id"] for x in got)


def test_mentor_comment_on_section_or_whole_document(app_client):
    rid, _ = reviewed(app_client)
    url = f"/api/reviews/{rid}/comments"
    r = app_client.post(url, json={"section_id": "S1", "problem": "Explain the old system."})
    assert r.status_code == 201 and r.json()["section_id"] == "S1" and not r.json()["anchored"]
    r = app_client.post(url, json={"problem": "Add a references section."})
    assert r.status_code == 201 and r.json()["section_id"] is None


def test_mentor_comment_validation(app_client):
    rid, _ = reviewed(app_client)
    url = f"/api/reviews/{rid}/comments"
    r = app_client.post(url, json={"paragraph_id": "S9.P9", "problem": "x"})
    assert r.status_code == 422 and r.json()["code"] == "BAD_ANCHOR"
    r = app_client.post(url, json={"section_id": "S9", "problem": "x"})
    assert r.status_code == 422 and r.json()["code"] == "BAD_ANCHOR"
    r = app_client.post(url, json={"paragraph_id": "S2.P1", "quote": "not there", "problem": "x"})
    assert r.status_code == 422 and r.json()["code"] == "BAD_QUOTE"
    assert app_client.post(url, json={"problem": ""}).status_code == 422
    assert app_client.post("/api/reviews/999/comments", json={"problem": "x"}).status_code == 404


def test_only_mentor_comments_can_be_deleted(app_client):
    rid, cs = reviewed(app_client)
    r = app_client.delete(f"/api/comments/{ai(cs)[0]['id']}")
    assert r.status_code == 403 and r.json()["code"] == "NOT_DELETABLE"
    mine = app_client.post(f"/api/reviews/{rid}/comments", json={"problem": "x"}).json()
    assert app_client.delete(f"/api/comments/{mine['id']}").status_code == 204
    assert app_client.delete(f"/api/comments/{mine['id']}").status_code == 404


def test_bulk_status(app_client):
    rid, cs = reviewed(app_client)
    ids = [c["id"] for c in cs if c["severity"] == "high"]
    r = app_client.post(
        f"/api/reviews/{rid}/comments/bulk", json={"ids": ids, "status": "accepted"}
    )
    assert r.status_code == 200 and {c["status"] for c in r.json()} == {"accepted"}
    assert len(r.json()) == len(ids)
    bad = app_client.post(
        f"/api/reviews/{rid}/comments/bulk", json={"ids": [999], "status": "rejected"}
    )
    assert bad.status_code == 422 and bad.json()["code"] == "UNKNOWN_COMMENTS"
    edit = app_client.post(
        f"/api/reviews/{rid}/comments/bulk", json={"ids": ids, "status": "edited"}
    )
    assert edit.status_code == 422


def test_overall_feedback_needs_confirmed_comments(app_client):
    rid, _ = reviewed(app_client)
    r = app_client.post(f"/api/reviews/{rid}/overall")
    assert r.status_code == 409 and r.json()["code"] == "NO_CONFIRMED_COMMENTS"


def test_overall_feedback_uses_only_confirmed_comments(app_client, chat):
    rid, cs = reviewed(app_client)
    first, second = ai(cs)[:2]
    patch(app_client, first["id"], status="edited", mentor_text="MENTOR-EDITED TEXT")
    patch(app_client, second["id"], status="rejected")
    r = app_client.post(f"/api/reviews/{rid}/overall")
    assert r.status_code == 200
    assert r.json()["overall_feedback"] == "Good structure; add evidence to the results."
    [prompt] = chat.prompts
    assert "MENTOR-EDITED TEXT" in prompt and second["problem"] not in prompt
    assert "research_problem" in prompt  # summary included
    assert "1 Approach" in prompt and "S1.P1" not in prompt  # student never sees the labels
    got = app_client.get(f"/api/reviews/{rid}").json()
    assert got["overall_feedback"].startswith("Good structure")


def test_mentor_edits_overall_feedback(app_client):
    rid, _ = reviewed(app_client)
    r = app_client.put(f"/api/reviews/{rid}/overall", json={"text": "My own words."})
    assert r.status_code == 200 and r.json()["overall_feedback"] == "My own words."


def test_overall_and_export_while_running_409(app_client):
    rid, cs = reviewed(app_client)
    patch(app_client, ai(cs)[0]["id"], status="accepted")
    set_status(app_client, rid, "running")
    assert app_client.post(f"/api/reviews/{rid}/overall").json()["code"] == "REVIEW_RUNNING"
    r = app_client.get(f"/api/reviews/{rid}/export", params={"format": "md"})
    assert r.status_code == 409 and r.json()["code"] == "REVIEW_RUNNING"


def test_export_markdown_with_chinese_filename(app_client):
    rid, cs = reviewed(app_client, name="实践报告.md")
    patch(app_client, ai(cs)[0]["id"], status="edited", mentor_text="Show the survey.")
    r = app_client.get(f"/api/reviews/{rid}/export", params={"format": "md"})
    assert r.status_code == 200 and r.headers["content-type"].startswith("text/markdown")
    assert f"filename*=UTF-8''{quote('实践报告-review.md')}" in r.headers["content-disposition"]
    assert "Show the survey." in r.text and ai(cs)[1]["problem"] not in r.text


def test_export_docx_of_a_docx_upload_has_word_comments(app_client):
    d = Docx()
    d.add_heading("1 Approach", 1)
    d.add_paragraph("We moved the order system to microservices. It is much better now.")
    buf = io.BytesIO()
    d.save(buf)
    doc_id = app_client.post(
        "/api/documents",
        files={"file": ("报告.docx", buf.getvalue())},
        data={"doc_type": "report"},
    ).json()["id"]
    rid = app_client.post("/api/reviews", json={"document_id": doc_id}).json()["id"]
    cs = app_client.get(f"/api/reviews/{rid}").json()["comments"]
    target = next(c for c in ai(cs) if c["paragraph_id"])
    patch(app_client, target["id"], status="accepted")
    r = app_client.get(f"/api/reviews/{rid}/export", params={"format": "docx"})
    assert r.status_code == 200 and "wordprocessingml" in r.headers["content-type"]
    out = Docx(io.BytesIO(r.content))
    assert len(list(out.comments)) == 1
    assert "microservices" in "\n".join(p.text for p in out.paragraphs)


def test_export_bad_format_422(app_client):
    rid, _ = reviewed(app_client)
    r = app_client.get(f"/api/reviews/{rid}/export", params={"format": "pdf"})
    assert r.status_code == 422
