"""Mentor-confirmed feedback → Markdown, or Word (comments on the student's own .docx)."""

import hashlib
import io
from datetime import UTC, datetime
from pathlib import Path

from app.models import Review, ReviewComment
from app.review import _loose
from app.schemas import Paragraph, ParsedDocument, StructuredSummary

EXPORTED = ("accepted", "edited")
EXCERPT = 200
_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"

LABELS = {
    "en": {
        "title": "Review feedback",
        "file": "File",
        "date": "Date",
        "model": "Model",
        "overall": "Overall feedback",
        "summary": "Summary",
        "comments": "Comments",
        "general": "General comments",
        "suggestion": "Suggestion",
        "steps": "Revision steps",
        "sources": "Sources",
        "figure": "Figure",
        "none": "No comments were confirmed.",
        "not_found": "not found",
        "table": ["Section", "Paragraph", "Comment", "Revision steps"],
        "fields": {
            "research_problem": "Research problem",
            "objectives": "Objectives",
            "technical_approach": "Technical approach",
            "main_conclusions": "Main conclusions",
            "innovations": "Innovations",
            "limitations": "Limitations",
        },
    },
    "zh": {
        "title": "评审意见",
        "file": "文件",
        "date": "日期",
        "model": "模型",
        "overall": "总体评价",
        "summary": "内容摘要",
        "comments": "具体意见",
        "general": "总体意见",
        "suggestion": "修改建议",
        "steps": "修改步骤",
        "sources": "依据",
        "figure": "图",
        "none": "没有已确认的意见。",
        "not_found": "未提及",
        "table": ["章节", "段落", "意见", "修改步骤"],
        "fields": {
            "research_problem": "研究问题",
            "objectives": "研究目标",
            "technical_approach": "技术路线",
            "main_conclusions": "主要结论",
            "innovations": "创新点",
            "limitations": "不足",
        },
    },
}


def exported(comments: list[ReviewComment]) -> list[ReviewComment]:
    """What the mentor confirmed: accepted, edited, and their own (saved as accepted)."""
    return [c for c in comments if c.status in EXPORTED]


def final_text(c: ReviewComment, lang: str = "en") -> str:
    if c.status == "edited" and c.mentor_text:
        return c.mentor_text
    tail = f"\n{LABELS[lang]['suggestion']}: {c.suggestion}" if c.suggestion else ""
    return c.problem + tail


def _ordered(parsed: ParsedDocument, comments: list[ReviewComment]):
    """[(section or None, [(paragraph or None, comment)])] in document order; None = whole doc."""
    pos = {pid: i for i, pid in enumerate(parsed.all_paragraph_ids())}
    groups = []
    for s in parsed.sections:
        cs = [c for c in comments if c.section_id == s.id]
        cs.sort(key=lambda c: pos.get(c.paragraph_id, -1))
        if cs:
            groups.append((s, [(parsed.paragraph(c.paragraph_id or ""), c) for c in cs]))
    known = {s.id for s in parsed.sections}
    rest = [c for c in comments if c.section_id not in known]
    if rest:
        groups.append((None, [(None, c) for c in rest]))
    return groups


def _excerpt(p: Paragraph, lang: str) -> str:
    if p.figure is not None:
        return f"{LABELS[lang]['figure']}: {p.figure.caption or p.figure.alt_text or p.id}"
    return p.text if len(p.text) <= EXCERPT else p.text[:EXCERPT] + "…"


def _summary_lines(review: Review, lang: str) -> list[tuple[str, str]]:
    if not review.summary:
        return []
    s = StructuredSummary(**review.summary)
    L = LABELS[lang]
    return [(label, getattr(s, f).text or L["not_found"]) for f, label in L["fields"].items()]


def _date(review: Review) -> str:
    return (review.finished_at or datetime.now(UTC)).strftime("%Y-%m-%d")


# ---- markdown -----------------------------------------------------------------------


def to_markdown(
    review: Review, filename: str, parsed: ParsedDocument, comments: list[ReviewComment]
) -> str:
    lang, cs = parsed.language, exported(comments)
    L = LABELS[lang]
    out = [
        f"# {L['title']} — {parsed.title or filename}",
        "",
        f"- {L['file']}: {filename}",
        f"- {L['date']}: {_date(review)}",
        f"- {L['model']}: {review.model}",
        "",
    ]
    if review.overall_feedback:
        out += [f"## {L['overall']}", "", review.overall_feedback, ""]
    if lines := _summary_lines(review, lang):
        out += [f"## {L['summary']}", ""] + [f"- **{k}:** {v}" for k, v in lines] + [""]
    out += [f"## {L['comments']} ({len(cs)})", ""]
    if not cs:
        out += [f"_{L['none']}_", ""]
    for section, items in _ordered(parsed, cs):
        out += [f"### {section.heading if section else L['general']}", ""]
        last = None
        for p, c in items:
            if p is not None and p.id != last:
                # ponytail: figure shown by caption only; embed the PNG (data URI) if asked
                out += [f"> [{p.id}] {_excerpt(p, lang)}", ""]
                last = p.id
            text = final_text(c, lang).replace("\n", "\n  ")
            out.append(f"- **[{c.severity} · {c.category}]** {text}")
            if c.revision_steps and c.status != "edited":
                out.append(f"  {L['steps']}:")
                out += [f"  {i}. {step}" for i, step in enumerate(c.revision_steps, 1)]
            if c.standard_refs:
                out.append(f"  {L['sources']}: " + "; ".join(c.standard_refs))
            out.append("")
    return "\n".join(out)


# ---- word ---------------------------------------------------------------------------


def _comment_body(c: ReviewComment, lang: str) -> str:
    text = final_text(c, lang)
    if c.revision_steps and c.status != "edited":
        steps = "\n".join(f"{i}. {s}" for i, s in enumerate(c.revision_steps, 1))
        text += f"\n{LABELS[lang]['steps']}:\n{steps}"
    return text


def _add_feedback(d, review: Review, lang: str) -> None:
    L = LABELS[lang]
    if review.overall_feedback:
        d.add_heading(L["overall"], 1)
        d.add_paragraph(review.overall_feedback)
    if lines := _summary_lines(review, lang):
        d.add_heading(L["summary"], 1)
        for k, v in lines:
            para = d.add_paragraph()
            para.add_run(f"{k}: ").bold = True
            para.add_run(v)


def to_docx(
    review: Review,
    filename: str,
    parsed: ParsedDocument,
    comments: list[ReviewComment],
    original: Path | None,
    figures_dir: Path,
) -> bytes:
    from docx import Document

    lang, cs = parsed.language, exported(comments)
    if original is not None:
        d = Document(str(original))
        general = _annotate(d, parsed, cs, lang)
        d.add_page_break()
        d.add_heading(LABELS[lang]["title"], 1)
        _add_feedback(d, review, lang)
        if general:
            d.add_heading(LABELS[lang]["general"], 1)
            for c in general:
                d.add_paragraph(f"[{c.severity} · {c.category}] {_comment_body(c, lang)}")
    else:
        d = _new_docx(review, filename, parsed, cs, figures_dir, lang)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _new_docx(review, filename, parsed, cs, figures_dir: Path, lang: str):
    from docx import Document
    from docx.shared import Inches

    L = LABELS[lang]
    d = Document()
    d.add_heading(f"{L['title']} — {parsed.title or filename}", 0)
    d.add_paragraph(f"{L['file']}: {filename}  ·  {L['date']}: {_date(review)}")
    _add_feedback(d, review, lang)
    d.add_heading(f"{L['comments']} ({len(cs)})", 1)
    if not cs:
        d.add_paragraph(L["none"])
        return d
    table = d.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for cell, name in zip(table.rows[0].cells, L["table"], strict=True):
        cell.text = name
    for section, items in _ordered(parsed, cs):
        for p, c in items:
            row = table.add_row().cells
            row[0].text = section.heading if section else L["general"]
            if p is not None:
                row[1].text = _excerpt(p, lang)
                png = figures_dir / p.figure.file if p.figure and p.figure.file else None
                if png and png.exists():
                    row[1].add_paragraph().add_run().add_picture(str(png), width=Inches(1.8))
            row[2].text = f"[{c.severity} · {c.category}] {final_text(c, lang)}"
            if c.status != "edited":
                row[3].text = "\n".join(f"{i}. {s}" for i, s in enumerate(c.revision_steps, 1))
    return d


def _annotate(d, parsed: ParsedDocument, cs: list[ReviewComment], lang: str):
    """Add a Word comment for each anchored comment; return the ones that found no place."""
    from docx.text.paragraph import Paragraph as DocxParagraph

    paras = [DocxParagraph(el, d._body) for el in d.element.body.iter(f"{_W}p")]
    loose = [_loose(p.text)[0] for p in paras]
    images = _image_paragraphs(d, paras)
    general = []
    for c in cs:
        i = _locate(c, parsed, loose, images)
        if i is None:
            general.append(c)
            continue
        # A paragraph whose text sits only in hyperlinks/fields has no runs to anchor on.
        order = [i, *range(i + 1, len(paras)), *range(i - 1, -1, -1)]
        j = next(k for k in order if paras[k].runs) if any(p.runs for p in paras) else None
        if j is None:
            general.append(c)
            continue
        d.add_comment(paras[j].runs, text=_comment_body(c, lang), author="Mentor", initials="M")
    return general


def _image_paragraphs(d, paras) -> dict[str, int]:
    """image_id (sha256[:32] of the image bytes, as the parser computes it) → paragraph index."""
    parts = d.part.related_parts
    out: dict[str, int] = {}
    for i, p in enumerate(paras):
        for blip in p._p.iter(f"{_A}blip"):
            if (part := parts.get(blip.get(f"{_R}embed"))) is not None:
                out.setdefault(hashlib.sha256(part.blob).hexdigest()[:32], i)
    return out


def _locate(c: ReviewComment, parsed, loose: list[str], images: dict[str, int]) -> int | None:
    """Index of the Word paragraph the comment belongs to: its paragraph, else its heading."""
    if c.section_id is None:
        return None
    section = next((s for s in parsed.sections if s.id == c.section_id), None)
    heading = _loose(section.heading)[0] if section else ""
    start = next((i for i, t in enumerate(loose) if heading and t == heading), None)
    p = parsed.paragraph(c.paragraph_id) if c.paragraph_id else None
    if p is not None and p.figure is not None and p.figure.image_id in images:
        return images[p.figure.image_id]
    if p is not None:
        # ponytail: text matching (quote, then the paragraph's opening words) from the section
        # heading on; mapping parser blocks to Word elements exactly would need parser changes
        text = p.figure.caption or "" if p.figure is not None else p.text
        keys = [_loose(k)[0] for k in (c.quote, text) if k]
        for key in (k[:40] for k in keys if k):
            for i in range(start or 0, len(loose)):
                t = loose[i]
                if key in t or (len(t) >= 4 and key.startswith(t) and p.kind == "table"):
                    return i
    return start
